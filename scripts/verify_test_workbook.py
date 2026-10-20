"""Read back the saved XLSX and independently check cached amounts and evidence."""
import json
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook

root=Path(__file__).resolve().parent.parent
payload=json.loads((root/"outputs/test_run_2025_records.json").read_text(encoding="utf-8"))
path=root/"outputs/financials_2025_test.xlsx"
wb=load_workbook(path,data_only=True)
formulas=load_workbook(path,data_only=False)
sheet=wb["Results"]
source=wb["Source data"]
assert wb.sheetnames==["Results","Source data"]
assert sheet.max_row==31 and sheet.max_column==9
rate=Decimal(payload["usdcad"])
previous=None
for i,record in enumerate(payload["records"],7):
    assert sheet.cell(i,1).value==i-6
    assert sheet.cell(i,2).value==record["company_name"]
    assert sheet.cell(i,8).value.date().isoformat()==record["fiscal_year_end"]
    assert sheet.cell(i,8).value.year==2025
    assert sheet.cell(i,7).value==("Yes" if record["converted_from_usd"] else "No")
    assert sheet.cell(i,9).value=="Ready"
    for col,key in enumerate(("revenue","net_income","total_assets"),4):
        item=record[key]
        expected=Decimal(str(item["value"]))*Decimal(item["unit_multiplier"])*(rate if item["currency"]=="USD" else Decimal(1))
        assert abs(Decimal(str(sheet.cell(i,col).value))-expected)<Decimal("0.001"),(record["company_name"],key)
        assert formulas["Results"].cell(i,col).data_type=="f"
    revenue=sheet.cell(i,4).value
    assert previous is None or previous>=revenue
    previous=revenue
    for col in (13,14,15,16,17,18,19,21,22):
        assert source.cell(i,col).value is not None,(i,col)
for tab in wb:
    for row in tab:
        assert all(cell.data_type!="e" for cell in row)
assert sum(sheet.cell(r,7).value=="Yes" for r in range(7,32))==16
print("Saved XLSX verified: 25 companies, 75 full CAD amounts, 25 ranks, 16 USD flags, fiscal ends, formulas, sources, explanations; no cached formula errors.")
