from decimal import Decimal
from pathlib import Path

from .validation import MANUAL, METRICS


def output_value(item: dict) -> int | float | str:
    """Output in thousands as in the example; no currency conversion or estimates."""
    if item["status"] != "ready" or item["value"] is None:
        return "Needs manual input"
    amount = Decimal(str(item["value"])) * Decimal(item["unit_multiplier"]) / Decimal(1000)
    return int(amount) if amount == amount.to_integral_value() else float(amount)


def write_workbook(records: list[dict], destination: Path, companies: list[dict]) -> None:
    """New workbook: 2025 values and evidence, including unresolved company rows.

    Callers must schema-validate and prepare records before calling this function.
    The example's 2024 figures are benchmarks only, outside the requested 2025 run.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    if not records:
        raise ValueError("No companies to export.")
    metadata = {company["company_id"]: company for company in companies}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Financials 2025"
    sheet.append(["Rank (reference)", "Company", "Stock symbol", "Fiscal year end",
                  "Assets '25 (000s)", "Net income '25 (000s)", "Revenue '25 (000s)",
                  "Assets currency", "Net income currency", "Revenue currency", "Status",
                  "Net income basis", "Net income derivation", "Review notes"])
    evidence = workbook.create_sheet("Evidence")
    evidence.append(["Company", "Fiscal year", "Metric", "Reported value", "Currency",
                     "Unit multiplier", "PDF file", "PDF page", "Printed page", "Source label",
                     "Source URL", "Publication date", "Auditor PDF page", "Original verified",
                     "Status", "Review note"])
    for record in records:
        company = metadata[record["company_id"]]
        sheet.append([company.get("rank"), record["company_name"], company.get("stock_symbol"),
                      record["fiscal_year_end"], *[output_value(record[key]) for key in
                      ("total_assets", "net_income", "revenue")],
                      *[record[key]["currency"] for key in ("total_assets", "net_income", "revenue")],
                      "Ready" if record["status"] == "ready" else "Needs manual input",
                      record["net_income_basis"], record["net_income_derivation"], record["review_notes"]])
        for metric in METRICS:
            item = record[metric]
            evidence.append([record["company_name"], record["fiscal_year"], metric, item["value"],
                             item["currency"], item["unit_multiplier"], record["pdf_filename"],
                             item["pdf_page"], item["printed_page"], item["source_label"],
                             record["source_url"], record["publication_date"],
                             record["auditor_evidence_pdf_page"], record["original_verified"],
                             item["status"], item["review_note"]])
    for tab in workbook:
        tab.freeze_panes = "D2"
        tab.auto_filter.ref = tab.dimensions
        tab.row_dimensions[1].height = 34
        for row in tab:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.value in (MANUAL, "Needs manual input"):
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
        for cell in tab[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
        for column in tab.columns:
            tab.column_dimensions[column[0].column_letter].width = 23
    sheet.column_dimensions["B"].width = 40
    sheet.column_dimensions["M"].width = 65
    sheet.column_dimensions["N"].width = 50
    evidence.column_dimensions["A"].width = 40
    evidence.column_dimensions["P"].width = 65
    for row in sheet.iter_rows(min_row=2, min_col=5, max_col=7):
        for cell in row:
            cell.number_format = '#,##0.###;(#,##0.###);0'
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("xb") as stream:
        workbook.save(stream)
